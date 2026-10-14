import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment
d=pd.read_pickle('/home/claude/work/post_det.pkl').sort_values(['anio','estrato'])
r=pd.read_pickle('/home/claude/work/post_res.pkl')
F='Arial'; blue=Font(name=F,color='0000FF'); blk=Font(name=F); bold=Font(name=F,bold=True); green=Font(name=F,color='008000')
hdr=PatternFill('solid',fgColor='D9E1F2')
wb=Workbook()
# ---------- Detalle ----------
ws=wb.active; ws.title='Detalle_estrato'
H=['Año','Estrato','n hogares ECV','% AC ECV (pond. FEX_C)','% ventilador ECV','% nevera ECV','Suscriptores SUI dic','Peso W','W·p AC','Var AC (W²·p(1−p)/n)','W·p ventilador','Var ventilador','W·p nevera','Var nevera']
ws.append(H)
for i,x in enumerate(d.itertuples(),start=2):
    ws.append([str(x.anio),'E'+x.estrato,x.n,x.p_ac,x.p_vent,x.p_nevera,x.suscriptores,
      f'=G{i}/SUMIFS($G:$G,$A:$A,A{i})',f'=H{i}*D{i}',f'=H{i}^2*D{i}*(1-D{i})/C{i}',
      f'=H{i}*E{i}',f'=H{i}^2*E{i}*(1-E{i})/C{i}',f'=H{i}*F{i}',f'=H{i}^2*F{i}*(1-F{i})/C{i}'])
    for c in 'ACDEFG': ws[f'{c}{i}'].font=blue if c in 'CDEFG' else blk
    for c in 'HIJKLMN': ws[f'{c}{i}'].font=blk
    for c in 'DEFHIKM': ws[f'{c}{i}'].number_format='0.0%'
    for c in 'JLN': ws[f'{c}{i}'].number_format='0.000000'
    ws[f'G{i}'].number_format='#,##0'
last=i
n=last+2
notes=['Fuentes: % por estrato = ECV DANE (P1077S11 AC, S12 ventilador, S2 nevera), Barranquilla, ponderado con FEX_C (ECV_2018_2025_Barranquilla_tenencia_por_estrato.csv).',
 'Suscriptores = SUI ele_com_096, Barranquilla urbano residencial, diciembre de cada año, suma de comercializadoras (SUI_Barranquilla_urbano_residencial_2015_2026.csv).',
 'Varianza: muestreo aleatorio simple por estrato, sin efecto de diseño. Cuando p = 0 o 1 (E6) el estrato aporta varianza 0; ver columna Agresti–Coull en Resumen.',
 'Se excluyen hogares sin estrato (grupo 0: rurales y sin dato), porque el SUI solo reparte suscriptores en E1–E6.']
for k,t in enumerate(notes): ws.cell(n+k,1,t).font=Font(name=F,italic=True)
# ---------- Resumen ----------
s=wb.create_sheet('Resumen_anual',0)
H2=['Año','n hogares E1–6','Suscriptores dic','% AC postestratificado','EE (pp)','IC95 ± (pp)','IC95 inf','IC95 sup','% AC pesos ECV (crudo)','Diferencia vs 2019 (pp)','z vs 2019','Efecto tasas vs 2019 (pp)','Efecto composición vs 2019 (pp)','IC95 ± Agresti–Coull (pp)','% ventilador postestr.','% nevera postestr.']
s.append(H2)
D="Detalle_estrato"
rows={}
for j,a in enumerate(r.index,start=2):
    rows[a]=j; A=f'$A{j}'
    s.append([str(a),f'=SUMIFS({D}!$C:$C,{D}!$A:$A,{A})',f'=SUMIFS({D}!$G:$G,{D}!$A:$A,{A})',
      f'=SUMIFS({D}!$I:$I,{D}!$A:$A,{A})',f'=SQRT(SUMIFS({D}!$J:$J,{D}!$A:$A,{A}))',f'=1.96*E{j}',f'=D{j}-F{j}',f'=D{j}+F{j}',
      r.loc[a,'pct_ac_pesosECV']/100,None,None,r.loc[a,'efecto_tasas_pp']/100,r.loc[a,'efecto_composicion_pp']/100,r.loc[a,'ic95_ac_AgCoull']/100,
      f'=SUMIFS({D}!$K:$K,{D}!$A:$A,{A})',f'=SUMIFS({D}!$M:$M,{D}!$A:$A,{A})'])
b=rows[2019]
for a,j in rows.items():
    s[f'J{j}']=f'=D{j}-$D${b}'
    s[f'K{j}']='' if a==2019 else f'=J{j}/SQRT(E{j}^2+$E${b}^2)'
    for c in 'DEFGHIJLMNOP': s[f'{c}{j}'].number_format='0.0%'
    s[f'K{j}'].number_format='0.0'; s[f'B{j}'].number_format='#,##0'; s[f'C{j}'].number_format='#,##0'
    for c in 'ABCDEFGHJKOP': s[f'{c}{j}'].font=blk
    for c in 'ILMN': s[f'{c}{j}'].font=blue
m=max(rows.values())+2
nt=['Lectura: % AC postestratificado = Σ W_e·p_e, con W_e = suscriptores SUI del estrato e en diciembre / total E1–E6 y p_e = tasa ECV dentro del estrato.',
 'EE = √Σ W_e²·p_e(1−p_e)/n_e (muestreo simple; el efecto de diseño de la ECV lo ampliaría). z vs 2019 trata ambos años como independientes.',
 'Efecto tasas = Σ W_2019·(p_año − p_2019); efecto composición = Σ (W_año − W_2019)·p_año. Suman la diferencia. Columnas I, L, M, N calculadas en Python (postestrat_anual.py).',
 '2019 con P1077S11 da 33,3 %; con P5523 (base ECV 2019 del informe) daba 33,5 %. 2020 fue en parte telefónica (pandemia). E5–E6 con n pequeños.',
 'Tendencia lineal ponderada (1/EE²) 2018–2025: +1,2 pp/año (EE 0,24); sin 2020–2021: +1,4 pp/año (EE 0,25).']
for k,t in enumerate(nt): s.cell(m+k,1,t).font=Font(name=F,italic=True)
for w in (s,ws):
    for c in w[1]: c.font=bold; c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='center')
    w.row_dimensions[1].height=45; w.freeze_panes='A2'
    for col in w.columns: w.column_dimensions[col[0].column_letter].width=14
wb.save('/mnt/user-data/outputs/Postestratificacion_AC_2018_2025.xlsx')
