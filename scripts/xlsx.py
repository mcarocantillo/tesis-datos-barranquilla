import pandas as pd, numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
O='/mnt/user-data/outputs/'
def rd(f,dec=','):
    d=pd.read_csv(O+f,sep=';',dtype=str,encoding='utf-8-sig',keep_default_na=False)
    for c in d.columns:
        s=d[c].str.strip()
        t=pd.to_numeric(s.str.replace(',','.') if dec==',' else s,errors='coerce')
        if t.notna().sum()>=0.9*(s!='').sum() and (s!='').sum()>0 and not c.startswith('DIRECTORIO'):
            d[c]=t
        else: d[c]=s.replace('',np.nan)
    return d
sheets={}
sheets['ECV2019_hogares']=rd('ECV2019_Barranquilla_consolidado.csv')
sheets['ECV2024_hogares']=rd('ECV2024_Barranquilla_hogares.csv',dec='.')
sheets['ENUT_personas']=rd('ENUT2024_25_Atlantico_cabecera_personas.csv')
sheets['ENUT_horas_tipo']=rd('ENUT2024_25_Atlantico_horas_por_tipo.csv')
sheets['GEIH2024_hogares']=rd('GEIH2024_Barranquilla_AM_hogares.csv')
sheets['GEIH2024_estrato']=rd('GEIH2024_Barranquilla_AM_resumen_estrato.csv')
sheets['SUI_residencial']=rd('SUI_Barranquilla_urbano_residencial_2019_2026.csv',dec='.')
sheets['SUI_no_residencial']=rd('SUI_Barranquilla_urbano_NO_residencial_2019_2026.csv',dec='.')
for k,v in sheets.items(): print(k,v.shape)
idx=pd.DataFrame([
 ['ECV2019_hogares','DANE, Encuesta de Calidad de Vida 2019 (cat. 678)','Barranquilla (08001)','Hogar',len(sheets['ECV2019_hogares']),'FEX_C','Vivienda, estrato, hogar, ingreso, AC (tipo y horas), ventilador (tipo y horas), nevera, electrodomésticos, tenencia, edades, ocupación (P6240)'],
 ['ECV2024_hogares','DANE, Encuesta de Calidad de Vida 2024 (cat. 861)','Barranquilla (08001)','Hogar',len(sheets['ECV2024_hogares']),'FEX_C','Vivienda, estrato, tenencia de AC, ventilador, nevera, lavadora; cuartos, sanitarios, personas, tenencia'],
 ['ENUT_personas','DANE, Encuesta Nacional de Uso del Tiempo 2024-2025 (cat. 910)','Atlántico cabecera (sin municipio)','Persona de 10+ años',len(sheets['ENUT_personas']),'FEX_C','Minutos por actividad el día de referencia, fecha y día de la semana, estrato, tipo de persona, lugar de trabajo, jornada de estudio'],
 ['ENUT_horas_tipo','Cálculo propio con ENUT 2024-2025','Atlántico cabecera','Resumen (tipo de persona × tipo de día)',len(sheets['ENUT_horas_tipo']),'ya ponderado','Horas promedio por actividad'],
 ['GEIH2024_hogares','DANE, Gran Encuesta Integrada de Hogares 2024, 12 meses (cat. 819)','Área metropolitana de Barranquilla (AREA 08)','Hogar',len(sheets['GEIH2024_hogares']),'FEX_ANUAL (= FEX_C18/12)','Vivienda, estrato, cuartos, valor vivienda, arriendo, composición del hogar, ocupados, trabajo en casa, horas'],
 ['GEIH2024_estrato','Cálculo propio con GEIH 2024','Área metropolitana de Barranquilla','Resumen por estrato',len(sheets['GEIH2024_estrato']),'ya ponderado','Promedios y porcentajes por estrato'],
 ['SUI_residencial','Superservicios, SUI reporte ele_com_096','Barranquilla urbano','Mes × estrato',len(sheets['SUI_residencial']),'no aplica (censal)','Suscriptores, consumo kWh, valor COP, kWh/suscriptor, COP/kWh; columna bandera = mes no válido'],
 ['SUI_no_residencial','Superservicios, SUI reporte ele_com_096','Barranquilla urbano','Mes',len(sheets['SUI_no_residencial']),'no aplica','Suscriptores y kWh industrial, comercial, oficial, otros (referencia; serie con saltos de reporte)'],
],columns=['Hoja','Fuente','Cobertura','Unidad (fila)','Filas','Factor de expansión','Contenido'])
dic=pd.DataFrame([
 ['Todas las encuestas','ESTRATO','Estrato para tarifa de energía','0 = sin estrato / recibo sin estrato; 1-6; 8 = planta eléctrica; 9 = no sabe'],
 ['ECV2019_hogares','TIPO_AC_P5523','Tipo de aire acondicionado','1 Mini-split; 2 Pared o ventana; 3 Central; vacío = sin AC'],
 ['ECV2019_hogares','HORAS_AC_P5524 / HORAS_VENTILADOR_P5526','Horas al día de uso (suma de todos los equipos)','Horas; puede superar 24'],
 ['ECV2019_hogares','VENT_*_P5525S1-S4','Tipo de ventilador: techo, pared, mesa, pie/pedestal','1 = sí, 0 = no'],
 ['ECV2019_hogares','TIENE_NEVERA / NEVERA_TAMANO_P5514 / NEVERA_ANTIGUEDAD_ANOS_P5515','Nevera inferida por respuesta; tamaño (códigos 1-3 sin etiqueta publicada); antigüedad en años (99 = no sabe)',''],
 ['ECV2019_hogares','TENENCIA_P5095','Tenencia de la vivienda','1 propia pagada; 2 propia pagando; 3 arriendo; 4 usufructo; 5 posesión sin título; 6 propiedad colectiva'],
 ['ECV2019_hogares','TIPO_VIVIENDA_P1070','Tipo de vivienda','1 casa; 2 apartamento; 3 cuarto; 5 otro'],
 ['ECV2019_hogares','N_*_P6240_*','Personas del hogar por condición de actividad (categorías estándar DANE, por confirmar)','1 trabajó; 2 buscó trabajo; 3 estudió; 4 oficios del hogar; 5 incapacitado; 6 otra'],
 ['ECV2024_hogares','TIENE_AC, TIENE_VENT, TIENE_NEVERA, TIENE_LAVADORA, TIENE_TV_LED, TIENE_CALENTADOR','Tenencia (P1077S11, S12, S2, S1, S8, S6)','1 = sí, 0 = no'],
 ['ECV2024_hogares','TIPO_VIVIENDA','Tipo de vivienda (P1070)','1 casa; 2 apartamento; 3 cuarto'],
 ['ECV2024_hogares','TENENCIA','Tenencia (P5095)','como 2019'],
 ['ENUT_personas','MIN_*','Minutos del día de referencia por actividad','Duraciones, sin hora de inicio/fin'],
 ['ENUT_personas','TIPO','Tipo de persona','Ocupado; Estudiante; Oficios del hogar; Mayor 65+ no ocupado; Busca trabajo; Otro / incapacitado'],
 ['ENUT_personas','LUGAR_TRABAJO','Dónde trabaja (P6880)','1 local fijo; 2 en esta vivienda; 3 otras viviendas; 4 kiosco; 5 vehículo; 6 puerta a puerta; 7 calle; ...'],
 ['ENUT_personas','JORNADA_ESTUDIO','Jornada (P1101)','1 mañana; 2 tarde; 3 noche; 4 única; 5 fines de semana'],
 ['ENUT_personas','MIN_FUERA_TRABAJO','Trabajo fuera de la vivienda + viajes (ida y vuelta)','minutos'],
 ['GEIH2024_hogares','TIPO_VIVIENDA','P4000','1 casa; 2 apartamento; 3 cuarto en inquilinato; 4 cuarto en otra estructura; 5 indígena; 6 otra'],
 ['GEIH2024_hogares','PAREDES','P4010','1 ladrillo/bloque; 2 madera pulida; 3 adobe; 4 bahareque; 5 madera burda; ... 9 sin paredes'],
 ['GEIH2024_hogares','PISOS','P4020','1 tierra; 2 cemento; 3 madera burda; 4 baldosín/cerámica; 5 mármol; 6 madera pulida; 7 alfombra'],
 ['GEIH2024_hogares','TENENCIA','P5090','1 propia pagada; 2 propia pagando; 3 arriendo; 4 usufructo; 5 posesión sin título; 6 otra'],
 ['GEIH2024_hogares','VALOR_VIVIENDA / ARRIENDO_*','P5110 precio mínimo de venta; P5130 arriendo estimado; P5140 arriendo pagado','COP'],
 ['GEIH2024_hogares','N_TRABAJA_EN_CASA','Ocupados del hogar que trabajan en esta vivienda (P6880 = 1)','personas'],
 ['SUI_residencial','bandera','Calidad del dato del mes','consumo_atipico, valor_negativo, valor_atipico; vacío = válido'],
],columns=['Hoja','Variable','Descripción','Códigos / unidad'])
path=O+'Bases_datos_tesis_Barranquilla.xlsx'
with pd.ExcelWriter(path,engine='openpyxl') as w:
    idx.to_excel(w,sheet_name='Indice',index=False)
    for k,v in sheets.items(): v.to_excel(w,sheet_name=k,index=False)
    dic.to_excel(w,sheet_name='Diccionario',index=False)
wb=load_workbook(path)
hdr=PatternFill('solid',start_color='DCE6F1')
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row: c.font=Font(name='Arial',size=10)
    for c in ws[1]: c.font=Font(name='Arial',size=10,bold=True); c.fill=hdr; c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for i,col in enumerate(ws.iter_cols(min_row=1,max_row=min(ws.max_row,200)),1):
        L=max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(i)].width=min(max(10,L+2),60 if ws.title in('Indice','Diccionario') else 28)
    if ws.title in ('Indice','Diccionario'):
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
wb.save(path); print('ok')
