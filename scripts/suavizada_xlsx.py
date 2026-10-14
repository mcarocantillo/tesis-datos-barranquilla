import pandas as pd, numpy as np, pickle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
S=pickle.load(open('/home/claude/work/suav.pkl','rb')); I=pickle.load(open('/home/claude/work/imput_anual.pkl','rb'))
F='Arial'; blue=Font(name=F,color='0000FF'); blk=Font(name=F); bold=Font(name=F,bold=True); it=Font(name=F,italic=True); grn=Font(name=F,color='008000')
hdr=PatternFill('solid',fgColor='D9E1F2')
def notes(ws,start,lines):
    for k,t in enumerate(lines): ws.cell(start+k,1,t).font=it
wb=Workbook()
# 1 Ajuste
a=wb.active; a.title='Ajuste_por_estrato'
a.append(['Estrato','b0 (logit en 2021)','b1 (cambio del logit por año)','EE de b1','p de b1','Sobredispersión (escala)','Olas usadas','Hogares en el ajuste'])
for r in S['F'].itertuples(): a.append([r.estrato,r.b0,r.b1,r.se_b1,r.p_b1,r.sobredispersion,r.olas,r.n_total])
for row in a.iter_rows(min_row=2):
    for cc in row[1:]: cc.font=blue
    for cc in row[1:6]: cc.number_format='0.000'
notes(a,a.max_row+2,['Modelo por estrato: logit(p) = b0 + b1·(año − 2021). GLM binomial con las tasas ponderadas de la ECV (P1077S11) y var_weights = n hogares.',
 'Olas: 2018, 2019, 2021, 2022, 2023, 2024, 2025. 2020 se excluye (encuesta en parte telefónica; año excluido de la calibración en la propuesta).',
 'Sobredispersión = estadístico de Pearson / gl (mínimo 1): cuánto varían las olas más de lo que explica el muestreo. Infla la varianza de cada año.',
 'Valores calculados en Python (serie_suavizada.py).'])
# 2 Serie por estrato
s=wb.create_sheet('Serie_estrato')
s.append(['Estrato','Año','Tipo de dato','n ECV','p observada (ECV)','p suavizada (media)','DE','IC90 inf','IC90 sup','Beta alfa','Beta beta','Control: media Beta'])
R=S['R'].sort_values(['estrato','anio'])
for i,r in enumerate(R.itertuples(),start=2):
    s.append([r.estrato,str(r.anio),r.tipo,None if np.isnan(r.n) else int(r.n),None if np.isnan(r.p_obs) else r.p_obs,r.p_suav,r.sd,r.ic90_inf,r.ic90_sup,
      f'=F{i}*(F{i}*(1-F{i})/G{i}^2-1)',f'=(1-F{i})*(F{i}*(1-F{i})/G{i}^2-1)',f'=J{i}/(J{i}+K{i})'])
    for col in 'DEFGHI': s[f'{col}{i}'].font=blue
    for col in 'EFGHIL': s[f'{col}{i}'].number_format='0.0%'
    for col in 'JK': s[f'{col}{i}'].number_format='0.00'
notes(s,i+2,['p suavizada y DE: del modelo de Ajuste_por_estrato (método delta, con sobredispersión). IC90 en escala logit.',
 'Beta por momentos: alfa = m·(m(1−m)/DE² − 1), beta = (1 − m)·(m(1−m)/DE² − 1). Es el prior de tenencia de cada estrato y año.',
 '2015-2017 y 2026-2027 son extrapolaciones de la tendencia (no hay ECV de Barranquilla para 2015-2017): su incertidumbre es mayor.'])
# 3 Ciudad
c=wb.create_sheet('Ciudad_postestratificada')
E=['E1','E2','E3','E4','E5','E6']
c.append(['Año']+[f'Peso SUI {e}' for e in E]+[f'p suavizada {e}' for e in E]+['% AC ciudad (suavizada)','DE (pp)','% AC ciudad (sin suavizar)'])
P=R.pivot(index='anio',columns='estrato',values='p_suav'); SD=R.pivot(index='anio',columns='estrato',values='sd')
for j,y in enumerate(P.index,start=2):
    raw=S['city'].loc[y,'pct_ac']
    c.append([str(y)]+[S['W'].loc[y,e] for e in E]+[P.loc[y,e] for e in E]+[f'=SUMPRODUCT(B{j}:G{j},H{j}:M{j})',
       '='+'SQRT('+'+'.join(f'({chr(66+k)}{j}*{SD.loc[y,E[k]]:.6f})^2' for k in range(6))+')*100',None if np.isnan(raw) else raw/100])
    for col in 'BCDEFGHIJKLMP': c[f'{col}{j}'].font=blue
    for col in 'BCDEFG': c[f'{col}{j}'].number_format='0.000'
    for col in 'HIJKLMNP': c[f'{col}{j}'].number_format='0.0%'
    c[f'O{j}'].number_format='0.0'
notes(c,j+2,['Pesos = suscriptores SUI de diciembre de cada año por estrato / total E1-E6. 2026 y 2027 usan los de diciembre de 2025.',
 'Sin suavizar = serie postestratificada de Postestratificacion_AC_2018_2025.xlsx, para comparar.'])
# 4 Imputación por año
m=wb.create_sheet('Imputacion_catastro')
m.append(['Año','Estrato','p objetivo (suavizada)','Intercepto a','b asbesto','b zinc/otro','P(AC|concreto)','P(AC|asbesto)','P(AC|zinc/otro)','Viviendas catastro',
 'Fracción en edificio 4+','Concreto en hasta 3 pisos','Asbesto en hasta 3 pisos','Zinc/otro en hasta 3 pisos','P(AC) edificio 4+','P(AC) hasta 3 pisos','P(AC) estrato (catastro)','Control: diferencia con objetivo','Viviendas con AC esperadas'])
Ix=I['I'].sort_values(['anio','estrato'])
for i,r in enumerate(Ix.itertuples(),start=2):
    m.append([str(r.anio),r.estrato,r.p_objetivo,r.intercepto,I['bt']['Asbesto'],I['bt']['Zinc/otro'],f'=1/(1+EXP(-D{i}))',f'=1/(1+EXP(-(D{i}+E{i})))',f'=1/(1+EXP(-(D{i}+F{i})))',
      r.viviendas,r.frac_edificio4,r.concreto_resto_hasta3,r.asbesto_hasta3,r.zinc_hasta3,f'=G{i}',f'=L{i}*G{i}+M{i}*H{i}+N{i}*I{i}',f'=K{i}*O{i}+(1-K{i})*P{i}',f'=Q{i}-C{i}',f'=J{i}*Q{i}'])
    for col in 'CDEFJKLMN': m[f'{col}{i}'].font=blue
    for col in 'CGHIKLMNOPQR': m[f'{col}{i}'].number_format='0.0%'
    for col in 'DEF': m[f'{col}{i}'].number_format='0.000'
    for col in 'JS': m[f'{col}{i}'].number_format='#,##0'
last=i
t=wb.create_sheet('Ciudad_catastro')
t.append(['Año','Viviendas','Viviendas con AC esperadas','% AC ciudad (catastro)'])
for j,y in enumerate(sorted(Ix.anio.unique()),start=2):
    t.append([str(y),f'=SUMIFS(Imputacion_catastro!$J$2:$J${last},Imputacion_catastro!$A$2:$A${last},A{j})',f'=SUMIFS(Imputacion_catastro!$S$2:$S${last},Imputacion_catastro!$A$2:$A${last},A{j})',f'=C{j}/B{j}'])
    t[f'B{j}'].number_format='#,##0'; t[f'C{j}'].number_format='#,##0'; t[f'D{j}'].number_format='0.0%'
notes(m,last+2,['Intercepto a calculado en Python (método de Brent) para que la media de las viviendas del catastro del estrato sea igual a la p objetivo (columna R = control, debe ser 0). Así E1 queda forzado a la ECV.',
 'b de techo: modelo M2 (techo + estrato) de la ECV 2024. Mezcla de techos: ECV 2024 para todos los años (simplificación: solo cambia el reparto dentro del estrato, no su total).',
 'Supuesto: vivienda en edificio de 4+ pisos (uso catastral "Apartamentos") = techo de concreto. Si cambia la p objetivo hay que recalcular a (imputacion_catastro_anual.py).'])
# 5 Posición
p=wb.create_sheet('Por_posicion')
p.append(['Posición frente a la cubierta','P(AC) media 2019','P(AC) media 2024'])
for k,r in I['pos'].iterrows(): p.append([k,r[2019],r[2024]])
for row in p.iter_rows(min_row=2):
    for cc in row[1:]: cc.number_format='0.0%'; cc.font=blue
for w in wb.worksheets:
    for cc in w[1]: cc.font=bold; cc.fill=hdr; cc.alignment=Alignment(wrap_text=True,vertical='center')
    w.row_dimensions[1].height=48; w.freeze_panes='A2'
    for col in w.columns: w.column_dimensions[col[0].column_letter].width=13
wb['Serie_estrato'].column_dimensions['C'].width=26; wb['Por_posicion'].column_dimensions['A'].width=34
wb.save('/mnt/user-data/outputs/Serie_suavizada_AC_2015_2027.xlsx')
R.to_csv('/mnt/user-data/outputs/Tenencia_AC_suavizada_estrato_2015_2027.csv',sep=';',index=False,encoding='utf-8-sig')
Ix[['anio','estrato','p_objetivo','pAC_edificio4','pAC_hasta3','pAC_estrato_catastro']].round(4).to_csv('/mnt/user-data/outputs/Catastro_PAC_por_anio_estrato_tipo.csv',sep=';',index=False,encoding='utf-8-sig')
